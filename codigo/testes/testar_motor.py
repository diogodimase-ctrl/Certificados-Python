"""
Testes do motor com arquivos de mentira (testes/amostras.py).

Rodar:  python testes/testar_motor.py            (rápido, sem PowerPoint)
        python testes/testar_motor.py --pdf      (também gera PDF de verdade)

Os valores esperados foram calculados à mão seguindo as regras do AHK v27.
"""

import json
import os
import re
import shutil
import sys
import tempfile
import time
import traceback

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from certificados import (caminhos, coordenadores, economia, gerador, historico, horas, painel,  # noqa: E402
                          planilhas, textos)
from testes import amostras  # noqa: E402

TESTES = []


def teste(fn):
    TESTES.append(fn)
    return fn


def igual(obtido, esperado, contexto=""):
    assert obtido == esperado, f"{contexto}\n  esperado: {esperado!r}\n  obtido:   {obtido!r}"


# ---------------------------------------------------------------------------
# Regras de texto e horas
# ---------------------------------------------------------------------------

@teste
def regras_de_texto():
    igual(textos.chave_nome("Dra. Márcia  Souza"), "marcia souza")
    igual(textos.chave_nome("Prof. Dr. João D'Ávila"), "joao d avila")
    igual(textos.limpar_nome_docente("Dra. Maria Exemplo (cardio)"), ("Maria Exemplo", "Dra. Maria Exemplo (cardio)"))
    igual(textos.limpar_nome_docente("Maria Exemplo M"), ("Maria Exemplo", "Maria Exemplo"))
    igual(textos.sanitizar_nome_arquivo('Pós: "A/B"?.'), "Pós AB")
    igual(textos.mes_para_pasta("Abril26"), "04.26")
    igual(textos.mes_para_pasta("Março/2026"), "03.26")
    igual(textos.mes_para_pasta("Planilha1"), "Planilha1")
    igual(textos.semestre_do_mes("04.26"), "2026-1")
    igual(textos.semestre_do_mes("10.26"), "2026-2")
    igual(textos.semestre_do_mes("02.27"), "2026-2")
    igual(textos.sigla_unidade("2026-A-PTA-1S"), "PTA")
    igual(textos.separar_coordenadores("A, B / C\\D"), ["A", "B", "C", "D"])
    igual(textos.juntar_nomes(["A", "B", "C"]), "A, B ou C")
    igual(textos.email_do_docente(["clarice.lima@x.com", "paulo@x.com"], "Clarice Lima"), "clarice.lima@x.com")
    igual(textos.extrair_email("Ricardo <ricardo.alves@x.com>"), "ricardo.alves@x.com")
    assert "2026-1s" in textos.variacoes_semestre("2026-5")


@teste
def regras_de_horas():
    igual(horas.formatar_horas("08:00 - 12:00"), "4 horas")
    igual(horas.formatar_horas("13:00 às 15:30"), "2 horas e 30 minutos")
    igual(horas.formatar_horas("19h às 22h"), "3 horas")
    igual(horas.formatar_horas("4:00"), "4 horas")
    igual(horas.formatar_horas("2h30"), "2 horas e 30 minutos")
    igual(horas.formatar_horas("1 hora"), "1 hora")
    igual(horas.formatar_horas("45min"), "45 minutos")
    assert horas.formatar_horas("manhã").startswith("Horario invalido")
    igual(horas.horas_de_valor_excel(0.125), "3 horas")
    igual(horas.horas_de_valor_excel(0.104166667), "2h30")
    igual(horas.formatar_duracao_digitada("2,30"), "2 horas e 30 minutos")
    igual(horas.formatar_duracao_digitada("1"), "1 hora")
    igual(horas.corrigir_data_digitada("5/3/26"), "05/03/2026")
    igual(horas.serial_para_data(46118), "06/04/2026")


# ---------------------------------------------------------------------------
# Planilhas, coordenadores e caminhos (ambiente de mentira)
# ---------------------------------------------------------------------------

def _ambiente():
    raiz = tempfile.mkdtemp(prefix="teste_cert_")
    return raiz, amostras.criar_ambiente(raiz)


@teste
def leitura_da_mestre():
    raiz, locais = _ambiente()
    try:
        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        igual([c.posgrad for c in cursos], ["Cardiologia Clínica", "Psiquiatria Adulto", "Curso Sem Planilha"])
        igual(cursos[0].sigla_unidade, "PTA")
        igual(cursos[0].unidade_exibida, "PTA (2026-A-PTA-1S)")
        igual(planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "não existe"), None)
        igual(planilhas.listar_abas(locais.planilha_mestre), ["2026-5", "2026-6"])
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def leitura_da_planilha_do_curso():
    raiz, locais = _ambiente()
    try:
        arq = caminhos.localizar_planilha_do_curso(raiz, "Cardiologia Clínica", "2026-A-PTA-1S", "2026-5")
        assert arq.endswith("Planejamento Cardiologia Clinica PTA 2026-1S.xlsx"), arq
        igual(caminhos.localizar_planilha_do_curso(raiz, "Curso Sem Planilha", "2026-A-RJ-1S", "2026-5"), "")

        dados = planilhas.ler_dados_da_planilha(arq, ["abril26", "Maio26"])
        igual(dados.nome_pos, "Pós-Graduação em Cardiologia Clínica")
        abril, maio = dados.abas
        igual(abril.colunas, ("E", "D", "F"), "colunas detectadas na linha 3")
        resumo = [(r.linha, r.nome, r.data_aula, r.horas_formatadas, r.email) for r in abril.registros]
        igual(resumo, [
            (5, "Maria Exemplo", "06/04/2026", "4 horas", "maria.exemplo@exemplo.com"),
            (6, "João Teste", "06/04/2026", "2 horas e 30 minutos", "joao@exemplo.com"),
            # Mesma quantidade de nomes e e-mails: pareados NA ORDEM (regra do AHK).
            (8, "Paulo Coelho", "07/04/2026", "4 horas", "clarice.lima@exemplo.com"),
            (8, "Clarice Lima", "07/04/2026", "4 horas", "paulo.coelho@exemplo.com"),
            # Data exibida como "8-abr": vem do número do Excel (o AHK daria 09/04).
            (9, "Maria Exemplo", "08/04/2026", "3 horas", "sem email"),
            (10, "Pedro Sem Aula", "09/04/2026", "2 horas e 30 minutos", "pedro@exemplo.com"),
        ])
        r5 = abril.registros[0]
        assert r5.mesclado and r5.suspeita_nome_com_extra and not r5.tem_erro
        assert abril.registros[-1].erro_curso
        igual([(r.nome, r.horas_formatadas, r.email) for r in maio.registros],
              [("Ricardo Alves", "3 horas", "ricardo.alves@exemplo.com")])
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def coordenadores_e_cadastro():
    raiz, locais = _ambiente()
    try:
        cad = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
        ana = coordenadores.resolver("Dra. Ana Beatriz Souza", cad, log=lambda m: None)
        igual((ana.nome, ana.titulo, ana.origem, ana.pendente), ("Dra. Ana Beatriz Souza", "Dra.", "pelo nome", ""))
        igual(ana.assina_como, "Dra. Ana Beatriz Souza")       # sem "Dra. Dra."
        igual(textos.sem_titulo("Prof. Dr. João Silva"), "João Silva")
        igual(ana.cargo, "Coordenadora do Curso")
        psi = coordenadores.resolver("Carlos Mendes / Paula Ribeiro", cad, log=lambda m: None)
        igual((psi.nome, psi.pendente), ("Paula Ribeiro", "titulo"))
        coordenadores.definir_titulo([ana, psi], cad, "Paula Ribeiro", "dra")
        igual((psi.titulo, psi.pendente), ("Dra.", ""))
        sem = coordenadores.resolver("Fulano Sem Assinatura", cad, log=lambda m: None)
        igual(sem.pendente, "assinatura")
        img = os.path.join(locais.assinaturas, "Ana Beatriz Souza.png")
        coordenadores.definir_assinatura(sem, cad, img)
        igual((sem.nome, sem.pendente), ("Fulano Sem Assinatura", "titulo"))
        # O cadastro foi gravado no formato do AHK e é lido de volta.
        cad2 = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
        igual(cad2.obter("paula ribeiro")["titulo"], "Dra.")
        igual(cad2.obter("Fulano Sem Assinatura")["assinatura"], "Ana Beatriz Souza.png")
        with open(cad.arquivo, encoding="utf-8-sig") as f:
            assert "Paula Ribeiro | Dra. | " in f.read()
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


# ---------------------------------------------------------------------------
# Geração completa
# ---------------------------------------------------------------------------

class ConversorFalso:
    """Não gera PDF de verdade: só copia o .pptx com nome .pdf."""
    nome = "falso"

    def converter(self, pptx, pdf):
        shutil.copy2(pptx, pdf)


def _lote(locais, raiz):
    cad = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
    curso = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")[0]
    c = gerador.CursoNoLote(curso)
    c.coordenacao = coordenadores.resolver(curso.coordenacao, cad, log=lambda m: None)
    c.planilha = caminhos.localizar_planilha_do_curso(raiz, curso.posgrad, curso.unidade, "2026-5")
    c.abas = ["Abril26"]
    c.dados = planilhas.ler_dados_da_planilha(c.planilha, c.abas)
    return c


def _gerar(modo_nome, conversor, raiz, locais):
    modo = gerador.MODOS[modo_nome]
    modo = gerador.Modo(**{**modo.__dict__, "pasta_local": os.path.join(raiz, "saida_" + modo_nome)})
    c = _lote(locais, raiz)
    problemas, _, _ = gerador.validar([c])
    for p in problemas:          # igual ao "Pular" da tela de problemas
        p.registro.pular = p.registro.tem_erro
    logs = []
    log = gerador.Log(modo_nome, pasta=os.path.join(raiz, "logs"))
    rel = gerador.gerar([c], locais, locais.modelo, "20/10/2026", modo, conversor, log,
                        email_teste="eu@teste.com")
    return c, rel, modo


@teste
def geracao_oficial():
    raiz, locais = _ambiente()
    try:
        c, rel, modo = _gerar("oficial", ConversorFalso(), raiz, locais)
        assert not rel.falhas, rel.falhas
        igual(rel.total_gerados, 5)
        envio = os.path.join(locais.declaracoes, "Para Enviar", "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        nomes = sorted(os.listdir(envio))
        igual(nomes, ["Clarice Lima (07.04.2026).pdf", "João Teste (06.04.2026).pdf",
                      "Maria Exemplo (06.04.2026).pdf", "Maria Exemplo (08.04.2026).pdf",
                      "Paulo Coelho (07.04.2026).pdf", "manifesto.txt"])
        with open(os.path.join(envio, "manifesto.txt"), encoding="utf-8") as f:
            m = json.load(f)
        igual(list(m)[:4], ["versaoManifesto", "loteId", "dataGeracao", "curso"])
        igual((m["curso"], m["local"], m["nomeAba"], m["semestre"], m["pastaMes"]),
              ("Cardiologia Clínica", "PTA", "Abril26", "2026-5", "04.26"))
        igual((m["pastaDestino"], m["pastaPos"]),
              ("2026-5/PTA/Cardiologia Clínica (2026-A-PTA-1S)/04.26", "Cardiologia Clínica (2026-A-PTA-1S)"))
        igual(len(m["arquivosDoLote"]), 5)
        maria = next(d for d in m["certificados"] if d["docente"] == "Maria Exemplo")
        igual([a["nomeArquivo"] for a in maria["arquivos"]],
              ["Maria Exemplo (06.04.2026).pdf", "Maria Exemplo (08.04.2026).pdf"])
        igual(maria["email"], "maria.exemplo@exemplo.com")
        docentes_local = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "Docentes")
        assert os.path.isdir(os.path.join(docentes_local, "Maria Exemplo"))
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def geracao_teste_e_simulacao():
    raiz, locais = _ambiente()
    try:
        _, rel, _ = _gerar("teste", ConversorFalso(), raiz, locais)
        envio = os.path.join(locais.declaracoes, "Para Enviar", "TESTE - Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        assert "TESTE - Maria Exemplo (06.04.2026).pdf" in os.listdir(envio)
        with open(os.path.join(envio, "manifesto.txt"), encoding="utf-8") as f:
            m = json.load(f)
        igual({d["email"] for d in m["certificados"]}, {"eu@teste.com"})
        assert m["pastaDestino"].startswith("TESTE/2026-5/PTA/")

        antes = set(os.listdir(os.path.join(locais.declaracoes, "Para Enviar")))
        _, rel, modo = _gerar("simulacao", ConversorFalso(), raiz, locais)
        igual(set(os.listdir(os.path.join(locais.declaracoes, "Para Enviar"))), antes, "simulação não pode tocar na nuvem")
        pasta = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        assert "manifesto.txt" in os.listdir(pasta)
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def certificado_preenchido():
    """O .pptx sai sem nenhum {MARCADOR}, com logo e assinatura trocados."""
    from pptx import Presentation
    from certificados import modelo
    raiz, locais = _ambiente()
    try:
        base = modelo.criar_modelo_base(locais.modelo, "Dra. Ana Beatriz Souza", "Coordenadora do Curso",
                                        "20/10/2026", os.path.join(locais.assinaturas, "Ana Beatriz Souza.png"),
                                        locais.logo)
        destino = os.path.join(raiz, "c.pptx")
        modelo.preencher_certificado(base, "Maria Exemplo", "Pós em Teste", "Anatomia", "06/04/2026", "4 horas", destino)
        pres = Presentation(destino)
        texto = "\n".join(f.text_frame.text for f in pres.slides[0].shapes if f.has_text_frame)
        assert "{" not in texto, texto
        assert "Declaramos que Maria Exemplo ministrou a aula Anatomia da Pós em Teste, em 06/04/2026, com carga horária de 4 horas." in texto
        assert "São Paulo, 20/10/2026" in texto
        nomes = [f.name for f in pres.slides[0].shapes]
        assert "ASSINATURA_REAL" in nomes and "LOGO_EINSTEIN" in nomes, nomes
        coord = next(f for f in pres.slides[0].shapes if "Dra. Ana" in (f.text_frame.text if f.has_text_frame else ""))
        assinatura = next(f for f in pres.slides[0].shapes if f.name == "ASSINATURA_REAL")
        igual(assinatura.left, coord.left)
        run = coord.text_frame.paragraphs[0].runs[0]
        igual((run.text, run.font.size.pt, run.font.bold), ("Dra. Ana Beatriz Souza", 10.0, False))
        run = coord.text_frame.paragraphs[1].runs[0]
        igual((run.text, run.font.size.pt, run.font.bold), ("Coordenadora do Curso", 10.0, True))
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def historico_de_envios():
    raiz, locais = _ambiente()
    try:
        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        cardio = cursos[0]
        igual(historico.situacao(locais, cursos)[cardio.linha], [])
        c, rel, _ = _gerar("oficial", ConversorFalso(), raiz, locais)
        assert historico.registrar(locais, rel, log=lambda m: None)
        # Gerado e ainda na Para Enviar: amarelo.
        meses = historico.situacao(locais, cursos)[cardio.linha]
        igual([(m.rotulo, m.pendente) for m in meses], [("Abr26", True)])
        # O fluxo arquivou e apagou a pasta da Para Enviar: verde, com a data do envio.
        destino = os.path.join(locais.declaracoes, "2026-5", "PTA", "Cardiologia Clínica", "04.26")
        shutil.move(os.path.join(locais.declaracoes, "Para Enviar", c.nome_com_unidade, "04.26"), destino)
        meses = historico.situacao(locais, cursos)[cardio.linha]
        igual([(m.rotulo, m.pendente) for m in meses], [("Abr26", False)])
        assert re.match(r"\d\d/\d\d$", meses[0].data), meses[0].data
        # Pasta arquivada sem histórico (envio do AHK) também conta; outro curso não é afetado.
        os.makedirs(os.path.join(locais.declaracoes, "2026-5", "PTA", "Cardiologia Clínica", "03.26"))
        with open(os.path.join(locais.declaracoes, "2026-5", "PTA", "Cardiologia Clínica", "03.26", "a.pdf"), "w"):
            pass
        sit = historico.situacao(locais, cursos)
        igual([m.rotulo for m in sit[cardio.linha]], ["Mar26", "Abr26"])
        igual(sit[cardio.linha][0].pdfs, 1)
        assert sit[cardio.linha][1].pdfs == rel.total_gerados > 0, (sit[cardio.linha][1].pdfs, rel.total_gerados)
        igual(sit[cursos[1].linha], [])
        igual(historico.meses_conhecidos(sit), ["04.26", "03.26"])
        igual(historico.rotulo_mes("10.26"), "Out26")
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def painel_da_tela_inicial():
    from datetime import date
    igual(painel.meses_do_semestre("2026.2"), ["09.26", "10.26", "11.26", "12.26", "01.27", "02.27"])
    igual(painel.meses_do_semestre("2026-5"), ["03.26", "04.26", "05.26", "06.26", "07.26", "08.26"])
    igual(painel.meses_do_semestre("Planilha1", hoje=date(2027, 1, 10))[0], "09.26")
    meses = painel.meses_do_semestre("2026.2")
    igual(painel.mes_padrao(meses, hoje=date(2026, 10, 22)), "10.26")
    igual(painel.mes_padrao(meses, hoje=date(2027, 5, 1)), "02.27")
    igual(painel.mes_padrao(meses, hoje=date(2026, 7, 1)), "09.26")
    igual((painel.curto("10.26"), painel.longo("01.27")), ("Out", "Janeiro/27"))
    assert painel.ja_passou("10.26", hoje=date(2026, 10, 1)) and not painel.ja_passou("11.26", hoje=date(2026, 10, 31))
    raiz, locais = _ambiente()
    try:
        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        cardio = cursos[0]
        planilha = caminhos.localizar_planilha_do_curso(locais.raiz_onedrive, cardio.posgrad, cardio.unidade, "2026-5")
        esperado = len(planilhas.ler_dados_da_planilha(planilha, ["Abril26"]).abas[0].registros)
        igual(planilhas.contar_certificados_do_mes(planilha, "04.26"), esperado)
        igual(planilhas.contar_certificados_do_mes(planilha, "06.26"), 0)
        cache = os.path.join(raiz, "contagem.json")
        contagem = painel.contar_previstos(locais, cursos, "2026-5", "04.26", arquivo=cache)
        igual(contagem[cardio.linha], esperado)
        assert os.path.isfile(cache)
        igual(painel.contar_previstos(locais, cursos, "2026-5", "04.26", arquivo=cache), contagem)   # do cache
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def semestre_da_turma_e_mestre_com_meses():
    igual([textos.semestre_da_turma("2026-1S"), textos.semestre_da_turma("2025-2S"),
           textos.semestre_da_turma("", "2026-A-CEP-1S"), textos.semestre_da_turma("Planilha1", "2026-B-RJ-2S"),
           textos.semestre_da_turma("2026-5"), textos.semestre_da_turma("x", "y")],
          ["2026-5", "2025-6", "2026-5", "2026-6", "2026-5", ""])
    from datetime import datetime as dt
    igual([planilhas.mes_do_cabecalho(v) for v in ("mar/26", "Outubro/2026", dt(2026, 11, 1), "TURMA", None)],
          ["03.26", "10.26", "11.26", "", ""])
    import openpyxl
    pasta = tempfile.mkdtemp()
    try:
        caminho = os.path.join(pasta, "mestre.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026-1S"
        ws.append(["Coordenação", "Pós-graduação", "TURMA/UNIDADE", "mar/26", "abr/26", "set/26"])
        ws.append(["Fulana", "Cardiologia", "2026-A-CEP-1S", "Enviado", "Enviado", None])
        ws.append(["Beltrano", "Medicina da Dor", "2026-A-RJ-1S", "Enviado", "sem aula", None])
        wb.save(caminho)
        mestre = planilhas.ler_mestre(caminho, "2026-1S")
        igual(mestre.meses, [("03.26", 4), ("04.26", 5), ("09.26", 6)])
        cardio, dor = mestre.cursos
        igual((cardio.turma, cardio.sigla_unidade, cardio.semestre, cardio.aba_mestre), ("", "CEP", "2026-5", "2026-1S"))
        assert cardio.enviado("04.26") and not dor.enviado("04.26") and not cardio.enviado("09.26")
        igual(dor.marcas, {"03.26": "Enviado", "04.26": "sem aula"})
        # Ciclo fechado em setembro para Cardiologia (fora da Para Enviar) e pendente para a Dor.
        sit = {cardio.linha: [historico.MesProcessado("09.26", "Set26", "", False, 8)],
               dor.linha: [historico.MesProcessado("09.26", "Set26", "", True, 3),
                           historico.MesProcessado("03.26", "Mar26", "", False, 3)]}
        igual([(c.posgrad, m, col) for c, m, col in historico.para_marcar_na_mestre(mestre.cursos, sit, mestre.meses)],
              [("Cardiologia", "09.26", 6)])
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


@teste
def organizar_pasta_enviados():
    from certificados import organizar
    decl = tempfile.mkdtemp()
    try:
        def pdf(*partes, conteudo="x"):
            os.makedirs(os.path.join(decl, *partes[:-1]), exist_ok=True)
            with open(os.path.join(decl, *partes), "w") as f:
                f.write(conteudo)
        pdf("Enviados", "Cardiologia (2026-A-CEP-1S)", "04.26", "Ana.pdf")
        pdf("Enviados", "Cardiologia (2026-A-CEP-1S)", "04.26", "Bia.pdf")
        pdf("Enviados", "Cardiologia (2026-A-CEP-1S)", "04.26", "manifesto.txt")
        pdf("Enviados", "Medicina da Dor (2025-B-RJ-2S) (1)", "10.25", "Caio.pdf")
        pdf("Enviados", "Sem Codigo", "04.26", "Davi.pdf")
        pdf("2026-5", "CEP", "Cardiologia", "04.26", "Ana.pdf")                    # já estava lá, igual
        pdf("2026-5", "CEP", "Cardiologia", "04.26", "Bia.pdf", conteudo="outro")  # já estava, diferente
        plano = organizar.planejar(decl)
        igual(sorted((os.path.relpath(m.destino, decl) if m.destino else "", m.pdfs, bool(m.problema)) for m in plano),
              [("", 1, True), (os.path.join("2025-6", "RJ", "Medicina da Dor (2025-B-RJ-2S)", "10.25"), 1, False),
               (os.path.join("2026-5", "CEP", "Cardiologia", "04.26"), 2, False)])
        movidos, erros = organizar.mover(plano, log=lambda m: None)
        igual((movidos, erros), (2, []))
        destino = os.path.join(decl, "2026-5", "CEP", "Cardiologia", "04.26")
        igual(sorted(os.listdir(destino)), ["Ana.pdf", "Bia (Enviados).pdf", "Bia.pdf"])
        assert os.path.isfile(os.path.join(decl, "2025-6", "RJ", "Medicina da Dor (2025-B-RJ-2S)", "10.25", "Caio.pdf"))
        # O que não é PDF fica; a pasta sem destino fica; a pasta vazia sai.
        assert os.path.isfile(os.path.join(decl, "Enviados", "Cardiologia (2026-A-CEP-1S)", "04.26", "manifesto.txt"))
        assert os.path.isfile(os.path.join(decl, "Enviados", "Sem Codigo", "04.26", "Davi.pdf"))
        assert not os.path.exists(os.path.join(decl, "Enviados", "Medicina da Dor (2025-B-RJ-2S) (1)"))
        # O histórico acha o que está dentro de Antigas também.
        os.makedirs(os.path.join(decl, "Antigas"))
        shutil.move(os.path.join(decl, "2025-6"), os.path.join(decl, "Antigas", "2025-6"))
        achados = historico.arquivados(type("Locais", (), {"declaracoes": decl})())
        assert "10.25" in achados[("rj", "medicina da dor", "2025-b-rj-2s")]
        assert "04.26" in achados[("cep", "cardiologia", "")]
    finally:
        shutil.rmtree(decl, ignore_errors=True)


@teste
def pasta_da_pos_sem_duplicar():
    decl = tempfile.mkdtemp()
    try:
        sigla = os.path.join(decl, "2026-5", "CEP")
        # 3) nada existe: cria no modelo "<Pós> (<código>)"
        igual(caminhos.pasta_da_pos(sigla, "Medicina da Dor", "2026-A-CEP-1S"), "Medicina da Dor (2026-A-CEP-1S)")
        # 2) só existe a pasta com o nome da pós: usa ela
        os.makedirs(os.path.join(sigla, "medicina da dor"))
        igual(caminhos.pasta_da_pos(sigla, "Medicina da Dor", "2026-A-CEP-1S"), "medicina da dor")
        # 1) existe com o código (acento/maiúscula diferentes): ela ganha da só-nome
        os.makedirs(os.path.join(sigla, "Medicina da DOR (2026-a-cep-1s)"))
        igual(caminhos.pasta_da_pos(sigla, "Medicina da Dor", "2026-A-CEP-1S"), "Medicina da DOR (2026-a-cep-1s)")
        # outra turma da mesma pós não pega a pasta com código de outra turma
        igual(caminhos.pasta_da_pos(sigla, "Medicina da Dor", "2026-B-CEP-1S"), "medicina da dor")
        igual(caminhos.separar_codigo("Enfermagem (UTI)"), ("Enfermagem (UTI)", ""))
        igual(caminhos.separar_codigo("Medicina da Dor (2026-A-CEP-1S)"), ("Medicina da Dor", "2026-A-CEP-1S"))

        # As barrinhas acham as duas formas de pasta.
        for pos, n in (("Medicina da Dor (2026-A-CEP-1S)", "09.26"), ("medicina da dor", "10.26"),
                       ("Medicina da Dor (2026-B-CEP-1S)", "11.26")):
            os.makedirs(os.path.join(sigla, pos, n), exist_ok=True)
            open(os.path.join(sigla, pos, n, "a.pdf"), "w").close()
        locais = type("Locais", (), {"declaracoes": decl})()
        curso = planilhas.Curso(2, "", "Medicina da Dor", "2026-A-CEP-1S", "CEP", "", "2026-1S")
        igual([m.pasta_mes for m in historico.situacao(locais, [curso])[2]], ["09.26", "10.26"])
    finally:
        shutil.rmtree(decl, ignore_errors=True)


@teste
def limpar_pastas_vazias_da_para_enviar():
    from certificados import organizar
    decl = tempfile.mkdtemp()
    try:
        pe = os.path.join(decl, "Para Enviar")
        os.makedirs(os.path.join(pe, "Vazia (2026-A-CEP-1S)"))
        os.makedirs(os.path.join(pe, "Com lote (2026-A-PTA-2S)", "10.26"))
        os.makedirs(os.path.join(pe, "Recém-criada (2026-A-RJ-1S)"))
        velho = time.time() - 3600
        for n in ("Vazia (2026-A-CEP-1S)", "Com lote (2026-A-PTA-2S)"):
            os.utime(os.path.join(pe, n), (velho, velho))
        igual(organizar.limpar_para_enviar(decl, log=lambda m: None), ["Vazia (2026-A-CEP-1S)"])
        igual(sorted(os.listdir(pe)), ["Com lote (2026-A-PTA-2S)", "Recém-criada (2026-A-RJ-1S)"])
    finally:
        shutil.rmtree(decl, ignore_errors=True)


@teste
def correcao_avulsa():
    from certificados import correcao
    raiz, locais = _ambiente()
    try:
        assert correcao.nome_confere("maria exemplo", "Dra. Maria  Exemplo")
        assert correcao.nome_confere("Maria", "Maria Exemplo") and not correcao.nome_confere("Mariana", "Maria Exemplo")
        assert correcao.data_confere("6/4", "06/04/2026") and correcao.data_confere("06.04.26", "06/04/2026")
        assert not correcao.data_confere("06/04/2025", "06/04/2026") and correcao.data_confere("", "06/04/2026")
        igual((correcao.horas_do_certificado("4"), correcao.horas_do_certificado("08:00 - 12:30"),
               correcao.horas_do_certificado("4 horas"), correcao.horas_do_certificado("meio período")),
              ("4 horas", "4 horas e 30 minutos", "4 horas", "meio período"))

        cursos = planilhas.ler_cursos_da_mestre(locais.planilha_mestre, "2026-5")
        guardadas = {}
        achados, sem, falhas = correcao.procurar(locais, cursos, "2026-5", "maria exemplo", guardadas=guardadas)
        igual(sorted(a.data for a in achados), ["06/04/2026", "08/04/2026"])
        a = achados[0]
        igual((a.curso, a.codigo, a.pasta_mes, a.docente), ("Cardiologia Clínica", "2026-A-PTA-1S", "04.26", "Maria Exemplo"))
        assert a.nome_pos and a.aula and a.horas, a
        igual(sem, ["Psiquiatria Adulto", "Curso Sem Planilha"])
        igual(falhas, [])
        achados, _, _ = correcao.procurar(locais, cursos, "2026-5", "maria exemplo", "8/4", guardadas=guardadas)
        igual([x.data for x in achados], ["08/04/2026"])

        cad = coordenadores.Cadastro(locais.assinaturas, log=lambda m: None)
        coord = coordenadores.resolver(cursos[0].coordenacao, cad, log=lambda m: None)
        p = correcao.Pedido("Maria Exemplo", a.nome_pos, "Aula corrigida", "6/4/26", "4", coord.nome, "Dra.",
                            "Coordenadora do Curso", coord.assinatura, "29/10/2026")
        igual(correcao.conferir(p), [])
        igual(correcao.conferir(correcao.Pedido("", "", "", "31/31", "", "", "", "", "", "")),
              ["Nome do docente em branco", "Pós-graduação em branco", "Aula / disciplina em branco",
               "Horas em branco", "Coordenador(a) em branco", "Data da aula inválida (use dd/mm/aaaa)",
               "Data de emissão inválida (use dd/mm/aaaa)", "Assinatura não escolhida"])
        pasta = os.path.join(raiz, "Correções")
        pdf, avisos = correcao.gerar(p, locais, locais.modelo, ConversorFalso(), pasta=pasta)
        igual(pdf, os.path.join(pasta, "04.26", "Maria Exemplo (06.04.2026).pdf"))
        from pptx import Presentation
        texto = " ".join(f.text_frame.text for f in Presentation(pdf).slides[0].shapes if f.has_text_frame)
        assert "Aula corrigida" in texto and "4 horas" in texto and "Maria Exemplo" in texto, texto
        pdf2, _ = correcao.gerar(p, locais, locais.modelo, ConversorFalso(), pasta=pasta)   # de novo: substitui
        igual((pdf2, os.listdir(os.path.dirname(pdf))), (pdf, ["Maria Exemplo (06.04.2026).pdf"]))
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


@teste
def tempo_economizado():
    pasta = tempfile.mkdtemp()
    try:
        arquivo = os.path.join(pasta, "tempo.json")
        igual(economia.resumo(arquivo).rodadas, 0)
        assert economia.registrar(23, 240, arquivo)
        assert economia.registrar(10, 100, arquivo)
        assert not economia.registrar(0, 50, arquivo)
        r = economia.resumo(arquivo)
        igual((r.rodadas, r.certificados, r.programa), (2, 33, 340.0))
        igual(r.manual, 33 * 693)
        igual(r.economizado, 33 * 693 - 340)
        igual((economia.duracao(3725), economia.duracao(95), economia.duracao(7)), ("1 h 02 min", "1 min 35 s", "7 s"))
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


def geracao_com_pdf_de_verdade(pasta_saida):
    """Gera os PDFs com o PowerPoint e confere o texto de dentro deles."""
    from pypdf import PdfReader
    from certificados.pdf import ConversorPowerPoint
    raiz, locais = _ambiente()
    try:
        with ConversorPowerPoint() as conv:
            _, rel, modo = _gerar("simulacao", conv, raiz, locais)
        pasta = os.path.join(modo.pasta_local, "Cardiologia Clínica (2026-A-PTA-1S)", "04.26")
        pdf = os.path.join(pasta, "Maria Exemplo (06.04.2026).pdf")
        texto = " ".join(PdfReader(pdf).pages[0].extract_text().split())
        assert "Maria Exemplo" in texto and "Anatomia do Coração" in texto and "4 horas" in texto, texto
        os.makedirs(pasta_saida, exist_ok=True)
        for n in os.listdir(pasta):
            shutil.copy2(os.path.join(pasta, n), pasta_saida)
        print("   PDFs copiados para", pasta_saida)
    finally:
        shutil.rmtree(raiz, ignore_errors=True)


if __name__ == "__main__":
    falhas = 0
    for fn in TESTES:
        try:
            fn()
            print(f"OK     {fn.__name__}")
        except Exception:
            falhas += 1
            print(f"FALHOU {fn.__name__}\n{traceback.format_exc()}")
    if "--pdf" in sys.argv:
        try:
            geracao_com_pdf_de_verdade(sys.argv[sys.argv.index("--pdf") + 1] if len(sys.argv) > sys.argv.index("--pdf") + 1
                                       else os.path.join(tempfile.gettempdir(), "certificados_pdf_teste"))
            print("OK     geracao_com_pdf_de_verdade")
        except Exception:
            falhas += 1
            print(f"FALHOU geracao_com_pdf_de_verdade\n{traceback.format_exc()}")
    print(f"\n{len(TESTES) + ('--pdf' in sys.argv) - falhas} de {len(TESTES) + ('--pdf' in sys.argv)} passaram")
    sys.exit(1 if falhas else 0)
